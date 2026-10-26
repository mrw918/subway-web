#!/usr/bin/env python3
"""根据当前 node-data 生成 7 份维护用 Excel（站点文案 + 课程）。"""

from __future__ import annotations

import json
import re
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "data" / "maintenance"
EXPORT_JS = ROOT / "scripts" / "export-node-data-json.js"


def safe_filename(title: str) -> str:
    name = re.sub(r'[\\/:*?"<>|]', "_", title.strip())
    return f"维护-{name}.xlsx"


def main() -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        print("需要 openpyxl：pip install openpyxl", file=sys.stderr)
        sys.exit(1)

    proc = subprocess.run(
        ["node", str(EXPORT_JS)],
        cwd=str(ROOT),
        capture_output=True,
        text=True,
        check=True,
    )
    routes = json.loads(proc.stdout)

    OUT_DIR.mkdir(parents=True, exist_ok=True)

    header_fill = PatternFill("solid", fgColor="F3F4F6")
    header_font = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    for route in routes:
        wb = Workbook()
        ws_site = wb.active
        ws_site.title = "站点文案"
        site_headers = ["站点名称", "站点说明", "资源类型"]
        ws_site.append(site_headers)
        for col in range(1, 4):
            cell = ws_site.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = header_font

        for node in route["nodes"]:
            ws_site.append(
                [
                    node["siteName"],
                    node["description"],
                    node["resourceTypes"],
                ]
            )

        for row in ws_site.iter_rows(min_row=2, max_row=ws_site.max_row):
            for cell in row:
                cell.alignment = wrap
        ws_site.column_dimensions["A"].width = 22
        ws_site.column_dimensions["B"].width = 72
        ws_site.column_dimensions["C"].width = 28

        ws_course = wb.create_sheet("课程")
        course_headers = ["站点名称", "课程名称", "课程类型", "链接", "课程说明"]
        ws_course.append(course_headers)
        for col in range(1, 6):
            cell = ws_course.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = header_font

        for node in route["nodes"]:
            courses = node["courses"]
            if not courses:
                ws_course.append([node["siteName"], "", "", "", ""])
                continue
            for course in courses:
                ws_course.append(
                    [
                        node["siteName"],
                        course["name"],
                        course["type"],
                        course["url"],
                        course["detail"],
                    ]
                )

        for row in ws_course.iter_rows(min_row=2, max_row=ws_course.max_row):
            for cell in row:
                cell.alignment = wrap
        ws_course.column_dimensions["A"].width = 22
        ws_course.column_dimensions["B"].width = 36
        ws_course.column_dimensions["C"].width = 12
        ws_course.column_dimensions["D"].width = 48
        ws_course.column_dimensions["E"].width = 72

        out_path = OUT_DIR / safe_filename(route["title"])
        wb.save(out_path)
        print(out_path.relative_to(ROOT))

    readme = OUT_DIR / "README.txt"
    readme.write_text(
        "维护说明\n"
        "========\n"
        "• 每人一条路线图对应一个 xlsx，站点名称请勿修改。\n"
        "• 「站点文案」：改站点说明、资源类型（顿号分隔，如 直播课、自学课）。\n"
        "• 「课程」：每个站点至少一行；无课则课名/链接/说明留空。\n"
        "  有课的站点：一门课一行（同一站点名称可重复多行）。\n"
        "• 底部标签切换「站点文案」与「课程」两个 Sheet。\n"
        "• 改完后将文件发给开发导入网站。\n"
        "• 重新生成模板（覆盖本地 xlsx）：python3 scripts/build-maintenance-xlsx.py\n",
        encoding="utf-8",
    )
    print(readme.relative_to(ROOT))


if __name__ == "__main__":
    main()
